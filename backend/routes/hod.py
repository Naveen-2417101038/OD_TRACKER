import os
import io
import re
from datetime import datetime
from werkzeug.utils import secure_filename
from flask import Blueprint, request, jsonify, send_from_directory, send_file, make_response, redirect
from sqlalchemy import or_, and_, func
from backend.database.sql_server import session_scope
from backend.models.db_models import User, Certificate, ODRequest, StudentAcademicRecord, to_iso
from backend.models.academic_record import AcademicRecordModel
from backend.services.od_eligibility import calculate_od_eligibility
from backend.models.od_request import ODRequestModel
from backend.models.user import UserModel
from backend.models.certificate import CertificateModel
from backend.routes.auth import role_required, current_user
from backend.config import Config

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

hod_bp = Blueprint('hod', __name__)

def is_hod_authorized_for_request(hod_user, req):
    """Verify that the request belongs to the HOD's department (or HOD has overall oversight)."""
    if not req or not hod_user:
        return False
    hod_dept = (hod_user.get('department') or '').strip().lower()
    req_dept = (req.get('department') or req.get('studentDepartment') or '').strip().lower()
    if hod_dept and hod_dept != 'all' and req_dept and hod_dept != req_dept:
        return False
    return True

def filter_od_requests_list(all_requests, filters):
    """Filter a list of OD requests in memory using provided filter criteria."""
    search_q = (filters.get('search') or '').strip().lower()
    event_name_q = (filters.get('event_name') or filters.get('eventName') or '').strip().lower()
    event_type_q = (filters.get('event_type') or filters.get('eventType') or '').strip().lower()
    status_q = (filters.get('status') or '').strip().lower()
    cert_status_q = (filters.get('certificate_status') or filters.get('certificateStatus') or '').strip().lower()
    student_name_q = (filters.get('student_name') or filters.get('studentName') or '').strip().lower()
    student_reg_q = (filters.get('student_reg_no') or filters.get('studentRegisterNo') or '').strip().lower()
    year_q = (filters.get('year') or '').strip().lower()
    section_q = (filters.get('section') or '').strip().lower()
    from_date_q = (filters.get('from_date') or filters.get('fromDate') or '').strip()
    to_date_q = (filters.get('to_date') or filters.get('toDate') or '').strip()

    filtered = []
    for r in all_requests:
        # General search match across name, reg no, event, id
        if search_q:
            matches_search = (
                search_q in str(r.get('id', '')).lower() or
                search_q in str(r.get('studentName', '')).lower() or
                search_q in str(r.get('studentRegisterNo', '')).lower() or
                search_q in str(r.get('eventName', '')).lower() or
                search_q in str(r.get('eventOrganizer', '')).lower() or
                search_q in str(r.get('venue', '')).lower()
            )
            if not matches_search:
                continue

        # Event Name filter
        if event_name_q:
            if event_name_q not in str(r.get('eventName', '')).lower():
                continue

        # Event Type filter
        if event_type_q and event_type_q != 'all':
            if event_type_q != str(r.get('eventType', '')).lower():
                continue

        # Status filter
        if status_q and status_q != 'all':
            req_status = str(r.get('status', '')).lower()
            if status_q == 'pending':
                # Any pending status (Mentor, Class Incharge, or HOD)
                if 'pending' not in req_status and req_status not in ['mentor approved', 'class incharge approved']:
                    continue
            elif status_q == 'pending_hod':
                if req_status != 'class incharge approved':
                    continue
            elif status_q == 'approved':
                if 'approved' not in req_status or req_status in ['mentor approved', 'class incharge approved']:
                    continue
            elif status_q == 'rejected':
                if 'rejected' not in req_status:
                    continue
            elif status_q != req_status:
                continue

        # Certificate Status filter
        if cert_status_q and cert_status_q != 'all':
            req_cert_status = str(r.get('certificateStatus', 'Not Uploaded')).lower()
            if cert_status_q == 'uploaded':
                if req_cert_status == 'not uploaded':
                    continue
            elif cert_status_q != req_cert_status:
                continue

        # Student Name filter
        if student_name_q:
            if student_name_q not in str(r.get('studentName', '')).lower():
                continue

        # Student Register Number filter
        if student_reg_q:
            if student_reg_q not in str(r.get('studentRegisterNo', '')).lower():
                continue

        # Year filter
        if year_q and year_q != 'all':
            if year_q not in str(r.get('studentYear', '')).lower():
                continue

        # Section filter
        if section_q and section_q != 'all':
            if section_q not in str(r.get('studentSection', '')).lower():
                continue

        # Date range filters (from_date / to_date)
        req_from = str(r.get('fromDate') or r.get('from_date') or '')
        req_to = str(r.get('toDate') or r.get('to_date') or req_from)

        if from_date_q and req_to:
            if req_to < from_date_q:
                continue

        if to_date_q and req_from:
            if req_from > to_date_q:
                continue

        filtered.append(r)

    return filtered

def generate_report_filename(filters):
    """Construct a clean, meaningful Excel filename based on applied filters."""
    now = datetime.now()
    month_year = now.strftime('%B_%Y') # e.g. September_2026
    
    event_name = (filters.get('event_name') or filters.get('eventName') or '').strip()
    event_type = (filters.get('event_type') or filters.get('eventType') or '').strip()
    status = (filters.get('status') or '').strip()

    prefix = "OD_Report"
    parts = [prefix]

    if event_name:
        # Sanitize event name into safe filename string
        clean_event = re.sub(r'[^a-zA-Z0-9_-]', '_', event_name)[:20].strip('_')
        if clean_event:
            parts.append(clean_event)
    elif event_type and event_type.lower() != 'all':
        parts.append(event_type.capitalize())

    if status and status.lower() != 'all':
        parts.append(status.capitalize().replace(' ', '_'))

    parts.append(month_year)
    filename = "_".join(parts) + ".xlsx"
    return filename

@hod_bp.route('/od-requests', methods=['GET'])
@role_required('HOD')
def get_hod_od_requests():
    """
    Retrieve all OD requests within the HOD's department oversight.
    Supports query parameter filters (search, event_type, status, date range, etc.).
    """
    hod_user = current_user
    all_requests = ODRequestModel.list_for_hod(hod_user)

    # Apply filters if provided in query parameters
    filters = request.args.to_dict()
    filtered_requests = filter_od_requests_list(all_requests, filters)

    pending_requests = [
        r for r in filtered_requests 
        if r.get('status') == 'Class Incharge Approved' and r.get('currentStage') == 'HOD'
    ]
    reviewed_requests = [
        r for r in filtered_requests 
        if not (r.get('status') == 'Class Incharge Approved' and r.get('currentStage') == 'HOD')
    ]
    approved_requests = [
        r for r in filtered_requests 
        if r.get('status') in ['Approved', 'HOD Approved', 'HOD Approved - Certificate Pending', 'Certificate Submitted']
    ]
    rejected_requests = [
        r for r in filtered_requests 
        if 'Rejected' in str(r.get('status'))
    ]

    status_filter = request.args.get('status', '').lower().strip()
    if status_filter == 'pending':
        return jsonify({
            'success': True,
            'requests': pending_requests,
            'count': len(pending_requests)
        }), 200
    elif status_filter in ['history', 'reviewed']:
        return jsonify({
            'success': True,
            'requests': reviewed_requests,
            'count': len(reviewed_requests)
        }), 200

    return jsonify({
        'success': True,
        'requests': filtered_requests,
        'all': filtered_requests,
        'pending': pending_requests,
        'history': reviewed_requests,
        'approved': approved_requests,
        'rejected': rejected_requests,
        'count': len(filtered_requests),
        'pendingCount': len(pending_requests)
    }), 200

@hod_bp.route('/od/export', methods=['GET'])
@hod_bp.route('/od-requests/export', methods=['GET'])
@role_required('HOD')
def export_hod_od_excel():
    """
    Export OD Request records as a professionally formatted Excel spreadsheet (.xlsx).
    Respects all active filter parameters (event name, event type, date range, status, student, etc.).
    """
    hod_user = current_user
    all_requests = ODRequestModel.list_for_hod(hod_user)

    # Read all filter parameters
    filters = request.args.to_dict()
    matching_records = filter_od_requests_list(all_requests, filters)

    # Create OpenPyXL workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    if not isinstance(ws, Worksheet):
        ws = wb.create_sheet("OD Requests Report")
    assert isinstance(ws, Worksheet)
    ws.title = "OD Requests Report"

    # Ensure grid lines are visible
    ws.sheet_view.showGridLines = True

    # Color definitions
    PRIMARY_COLOR = "1E1B4B"      # Dark Indigo
    HEADER_FILL = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    TITLE_FILL = PatternFill(start_color="311042", end_color="311042", fill_type="solid")
    TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    
    SUBTITLE_FILL = PatternFill(start_color="F5F3FF", end_color="F5F3FF", fill_type="solid")
    SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="4B5563")

    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    BORDER_THIN = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    # Status color fills
    STATUS_APPROVED_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    STATUS_APPROVED_FONT = Font(name="Calibri", size=10, bold=True, color="15803D")
    
    STATUS_REJECTED_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    STATUS_REJECTED_FONT = Font(name="Calibri", size=10, bold=True, color="B91C1C")
    
    STATUS_PENDING_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    STATUS_PENDING_FONT = Font(name="Calibri", size=10, bold=True, color="B45309")

    # Column Headers (25 comprehensive columns)
    headers = [
        "S.No",
        "Request ID",
        "Student Name",
        "Register Number",
        "Department",
        "Year",
        "Section",
        "Event / Program Name",
        "Event Type",
        "Event Date",
        "Event Venue / Organization",
        "OD Reason",
        "OD Request Date",
        "OD Letter / Document",
        "Faculty / Mentor Status",
        "Faculty / Mentor Remarks",
        "Class Incharge Status",
        "Class Incharge Remarks",
        "HOD Status",
        "HOD Remarks",
        "Overall OD Status",
        "Certificate Uploaded",
        "Certificate Verification Status",
        "Certificate Remarks",
        "Final Completion Status"
    ]

    total_cols = len(headers)

    # Title Banner (Row 1)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    title_cell = ws.cell(row=1, column=1, value="RAJALAKSHMI ENGINEERING COLLEGE (AUTONOMOUS) — ON-DUTY (OD) TRACKING REPORT")
    title_cell.font = TITLE_FONT
    title_cell.fill = TITLE_FILL
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Subtitle with metadata and filter summary (Row 2)
    dept_name = hod_user.get('department') or 'Computer Science and Design'
    gen_time = datetime.now().strftime('%d-%b-%Y %I:%M %p')
    filter_desc_items = []
    event_type_val = filters.get('event_type')
    if event_type_val and event_type_val.lower() != 'all':
        filter_desc_items.append(f"Event Type: {event_type_val}")
    event_name_val = filters.get('event_name')
    if event_name_val:
        filter_desc_items.append(f"Event: {event_name_val}")
    status_val = filters.get('status')
    if status_val and status_val.lower() != 'all':
        filter_desc_items.append(f"Status: {status_val}")
    if filters.get('from_date') or filters.get('to_date'):
        d_from = filters.get('from_date', 'Start')
        d_to = filters.get('to_date', 'End')
        filter_desc_items.append(f"Date Range: {d_from} to {d_to}")

    filter_desc = " | Filters: " + ", ".join(filter_desc_items) if filter_desc_items else " | Filters: All Records"

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    sub_cell = ws.cell(row=2, column=1, value=f"Department: {dept_name} | Generated On: {gen_time} | Total Matching Records: {len(matching_records)}{filter_desc}")
    sub_cell.font = SUBTITLE_FONT
    sub_cell.fill = SUBTITLE_FILL
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 24

    # Blank Separator (Row 3)
    ws.row_dimensions[3].height = 10

    # Header Row (Row 4)
    ws.row_dimensions[4].height = 28
    for col_num, header_text in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header_text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN

    # Data Rows (Starting at Row 5)
    row_idx = 5
    for idx, r in enumerate(matching_records, 1):
        ws.row_dimensions[row_idx].height = 22
        fill_to_use = ZEBRA_FILL if idx % 2 == 0 else WHITE_FILL

        # 1. Fetch real student/user record from database using student ID reference or register number
        student_id = r.get('studentId') or r.get('student_id')
        student_reg = r.get('studentRegisterNo') or r.get('student_reg_no')
        
        student_user = None
        if student_id:
            student_user = UserModel.get_by_id(student_id)
            if not student_user:
                student_user = UserModel.get_by_identifier(student_id)
        if not student_user and student_reg:
            student_user = UserModel.get_by_identifier(student_reg)

        # 2. Join actual student profile information
        if student_user:
            student_name = student_user.get('name') or 'Student Not Found'
            reg_no = student_user.get('identifier') or student_user.get('register_no') or student_user.get('student_reg_no') or 'N/A'
            department = student_user.get('department') or r.get('studentDepartment') or r.get('department') or 'N/A'
            year = student_user.get('year') or r.get('studentYear') or r.get('year') or 'N/A'
            section = student_user.get('section') or r.get('studentSection') or r.get('section') or 'N/A'
        else:
            # Handle edge case: Student not found / deleted from database
            student_name = "Student Not Found"
            reg_no = "N/A"
            department = r.get('studentDepartment') or r.get('department') or 'N/A'
            year = r.get('studentYear') or r.get('year') or 'N/A'
            section = r.get('studentSection') or r.get('section') or 'N/A'

        req_id = str(r.get('id') or r.get('requestId') or '')
        event_name = str(r.get('eventName') or r.get('event_name') or '')
        event_type = str(r.get('eventType') or r.get('event_type') or '')
        
        # Event date formatting
        from_d = str(r.get('fromDate') or r.get('from_date') or r.get('eventDate') or '')
        to_d = str(r.get('toDate') or r.get('to_date') or from_d)
        date_str = f"{from_d} to {to_d}" if from_d != to_d and to_d else from_d

        venue = f"{r.get('venue', '')} ({r.get('eventOrganizer', '')})" if r.get('venue') and r.get('eventOrganizer') else (r.get('venue') or r.get('eventOrganizer') or 'N/A')
        reason = str(r.get('reason') or '')
        created_at = str(r.get('createdAt') or r.get('created_at') or '')
        has_doc = "Attached" if (r.get('odLetterUrl') or r.get('documentUrl')) else "None"

        # Stages info
        stages = r.get('stages') or {}
        mentor_stage = stages.get('mentor') or {}
        ci_stage = stages.get('classIncharge') or {}
        hod_stage = stages.get('hod') or {}

        mentor_status = str(mentor_stage.get('status', 'Pending'))
        mentor_remarks = str(mentor_stage.get('feedback') or '')
        ci_status = str(ci_stage.get('status', 'Unreached'))
        ci_remarks = str(ci_stage.get('feedback') or '')
        hod_status = str(hod_stage.get('status', 'Unreached'))
        hod_remarks = str(r.get('remarks') if r.get('status') in ['Approved', 'HOD Approved', 'HOD Rejected'] else '')

        overall_status = str(r.get('status', 'Pending'))
        cert_status = str(r.get('certificateStatus', 'Not Uploaded'))
        cert_uploaded = "Yes" if cert_status != 'Not Uploaded' else "No"
        
        # Certificate remarks
        cert_record = CertificateModel.get_by_request_id(req_id) if req_id else None
        cert_remarks = str((cert_record.get('remarks') if cert_record else '') or r.get('certificateRemarks') or r.get('certificate_remarks') or '')

        if cert_status == 'Verified':
            final_status = "Completed & Verified"
        elif overall_status in ['Approved', 'HOD Approved', 'HOD Approved - Certificate Pending', 'Certificate Submitted']:
            final_status = "Approved - Awaiting Certificate" if cert_status in ['Not Uploaded', 'Pending Upload'] else f"Approved ({cert_status})"
        elif 'Rejected' in overall_status:
            final_status = "Rejected"
        else:
            final_status = f"In Progress ({r.get('currentStage', 'Mentor')})"

        # Debugging log output as required by specification
        debug_record = {
            "studentName": student_name,
            "registerNumber": reg_no,
            "department": department,
            "year": year,
            "section": section,
            "eventName": event_name,
            "overallStatus": overall_status
        }
        print(f"[EXPORT HOD EXCEL DEBUG] Row #{idx}: {debug_record}")

        row_values = [
            idx,
            req_id,
            student_name,
            reg_no,
            department,
            year,
            section,
            event_name,
            event_type,
            date_str,
            venue,
            reason,
            created_at,
            has_doc,
            mentor_status,
            mentor_remarks,
            ci_status,
            ci_remarks,
            hod_status,
            hod_remarks,
            overall_status,
            cert_uploaded,
            cert_status,
            cert_remarks,
            final_status
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = val
            cell.fill = fill_to_use
            cell.border = BORDER_THIN
            cell.font = Font(name="Calibri", size=10, color="1F2937")
            
            # Alignments
            if col_idx in [1, 2, 4, 6, 7, 9, 10, 13, 14, 15, 17, 19, 21, 22, 23, 25]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Status column highlights
            if col_idx == 21: # Overall Status
                if 'Approved' in overall_status:
                    cell.fill = STATUS_APPROVED_FILL
                    cell.font = STATUS_APPROVED_FONT
                elif 'Rejected' in overall_status:
                    cell.fill = STATUS_REJECTED_FILL
                    cell.font = STATUS_REJECTED_FONT
                else:
                    cell.fill = STATUS_PENDING_FILL
                    cell.font = STATUS_PENDING_FONT

            elif col_idx == 23: # Certificate status
                if cert_status == 'Verified':
                    cell.fill = STATUS_APPROVED_FILL
                    cell.font = STATUS_APPROVED_FONT
                elif cert_status == 'Pending Verification':
                    cell.fill = STATUS_PENDING_FILL
                    cell.font = STATUS_PENDING_FONT

            elif col_idx == 25: # Final completion status
                if 'Completed' in final_status or 'Approved' in final_status:
                    cell.fill = STATUS_APPROVED_FILL
                    cell.font = STATUS_APPROVED_FONT
                elif 'Rejected' in final_status:
                    cell.fill = STATUS_REJECTED_FILL
                    cell.font = STATUS_REJECTED_FONT

        row_idx += 1

    # If no records match
    if not matching_records:
        ws.row_dimensions[row_idx].height = 30
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=total_cols)
        empty_cell = ws.cell(row=row_idx, column=1, value="No OD requests match the currently selected filter criteria.")
        empty_cell.font = Font(name="Calibri", size=11, italic=True, color="64748B")
        empty_cell.alignment = Alignment(horizontal="center", vertical="center")
        empty_cell.fill = ZEBRA_FILL
        empty_cell.border = BORDER_THIN
        row_idx += 1

    # Summary Footer Row
    summary_row = row_idx + 1
    ws.row_dimensions[summary_row].height = 26
    ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=4)
    sum_title_cell = ws.cell(row=summary_row, column=1, value=f"EXECUTIVE SUMMARY: {len(matching_records)} Total Records")
    sum_title_cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sum_title_cell.fill = HEADER_FILL
    sum_title_cell.alignment = Alignment(horizontal="center", vertical="center")

    approved_cnt = len([r for r in matching_records if 'Approved' in str(r.get('status'))])
    rejected_cnt = len([r for r in matching_records if 'Rejected' in str(r.get('status'))])
    pending_cnt = len(matching_records) - approved_cnt - rejected_cnt
    verified_cert_cnt = len([r for r in matching_records if r.get('certificateStatus') == 'Verified'])

    ws.merge_cells(start_row=summary_row, start_column=5, end_row=summary_row, end_column=total_cols)
    sum_val_cell = ws.cell(row=summary_row, column=5, value=f"Approved: {approved_cnt} | Pending: {pending_cnt} | Rejected: {rejected_cnt} | Verified Certificates: {verified_cert_cnt}")
    sum_val_cell.font = Font(name="Calibri", size=10, bold=True, color="1E1B4B")
    sum_val_cell.fill = SUBTITLE_FILL
    sum_val_cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col_idx in range(1, total_cols + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for r_idx in range(4, row_idx):
            c_val = ws.cell(row=r_idx, column=col_idx).value
            if c_val is not None:
                max_len = max(max_len, len(str(c_val)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save to memory buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = generate_report_filename(filters)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )

@hod_bp.route('/od-requests/<request_id>', methods=['GET'])
@role_required('HOD')
def get_hod_od_request_detail(request_id):
    """Retrieve full details of a specific OD request for HOD executive review."""
    hod_user = current_user
    req = ODRequestModel.get_by_id(request_id)

    if not req:
        return jsonify({
            'success': False,
            'error': f"OD Request with ID '{request_id}' not found."
        }), 404

    if not is_hod_authorized_for_request(hod_user, req):
        return jsonify({
            'success': False,
            'error': "Access denied: This OD request belongs to another academic department."
        }), 403

    return jsonify({
        'success': True,
        'request': req
    }), 200

@hod_bp.route('/od-requests/<request_id>/approve', methods=['POST'])
@role_required('HOD')
def approve_od_request_by_hod(request_id):
    """
    Grant official executive approval to an OD application as HOD:
    - Validates request state is 'Class Incharge Approved'
    - Updates status to 'Approved'
    - Advances stage to 'Approved'
    - Logs decision in approvals and od_history collections
    - Notifies student that certificate upload is unlocked
    """
    hod_user = current_user
    data = request.get_json(silent=True) or {}
    remarks = (data.get('remarks') or data.get('comments') or '').strip()

    req = ODRequestModel.get_by_id(request_id)
    if not req:
        return jsonify({
            'success': False,
            'error': f"OD Request with ID '{request_id}' not found."
        }), 404

    if not is_hod_authorized_for_request(hod_user, req):
        return jsonify({
            'success': False,
            'error': "Access denied: You cannot sanction OD requests belonging to another department."
        }), 403

    updated_req, err = ODRequestModel.approve_by_hod(request_id, hod_user, remarks)
    if err:
        return jsonify({
            'success': False,
            'error': err
        }), 400

    return jsonify({
        'success': True,
        'message': f"OD Request {request_id} has been officially approved and sanctioned by HOD.",
        'request': updated_req,
        'status': updated_req.get('status') if updated_req else None,
        'currentStage': updated_req.get('currentStage') if updated_req else None
    }), 200

@hod_bp.route('/od-requests/<request_id>/reject', methods=['POST'])
@role_required('HOD')
def reject_od_request_by_hod(request_id):
    """
    Reject an OD application as HOD:
    - Requires mandatory rejection reason
    - Validates request state is 'Class Incharge Approved'
    - Updates status to 'HOD Rejected'
    - Logs decision in approvals and od_history collections
    """
    hod_user = current_user
    data = request.get_json(silent=True) or {}
    reason = (data.get('reason') or data.get('remarks') or data.get('comments') or '').strip()

    if not reason:
        return jsonify({
            'success': False,
            'error': "Rejection reason / remark is mandatory when rejecting an OD request."
        }), 400

    req = ODRequestModel.get_by_id(request_id)
    if not req:
        return jsonify({
            'success': False,
            'error': f"OD Request with ID '{request_id}' not found."
        }), 404

    if not is_hod_authorized_for_request(hod_user, req):
        return jsonify({
            'success': False,
            'error': "Access denied: You cannot reject OD requests belonging to another department."
        }), 403

    updated_req, err = ODRequestModel.reject_by_hod(request_id, hod_user, reason)
    if err:
        return jsonify({
            'success': False,
            'error': err
        }), 400

    return jsonify({
        'success': True,
        'message': f"OD Request {request_id} has been declined by HOD.",
        'request': updated_req,
        'status': updated_req.get('status') if updated_req else None,
        'rejectionReason': updated_req.get('rejectionReason') if updated_req else None
    }), 200

@hod_bp.route('/stats', methods=['GET'])
@role_required('HOD')
def get_hod_department_stats():
    """Compute live department analytics and statistics for HOD dashboard."""
    hod_user = current_user
    all_requests = ODRequestModel.list_for_hod(hod_user)

    total = len(all_requests)
    pending_hod = len([r for r in all_requests if r.get('status') == 'Class Incharge Approved'])
    pending_mentor = len([r for r in all_requests if r.get('status') == 'Pending'])
    pending_ci = len([r for r in all_requests if r.get('status') == 'Mentor Approved'])
    approved = len([r for r in all_requests if r.get('status') in ['Approved', 'HOD Approved', 'HOD Approved - Certificate Pending', 'Certificate Submitted']])
    rejected = len([r for r in all_requests if 'Rejected' in str(r.get('status'))])

    # Event distribution
    event_dist = {}
    for r in all_requests:
        ev_type = r.get('eventType') or r.get('event_type') or 'Other'
        event_dist[ev_type] = event_dist.get(ev_type, 0) + 1

    dept_lower = (hod_user.get('department') or '').strip().lower()
    with session_scope() as session:
        att_query = session.query(func.avg(User.attendance_percentage)).filter(
            or_(func.lower(User.role) == 'student', func.lower(User.sub_role) == 'student')
        )
        if dept_lower and dept_lower != 'all':
            if dept_lower in ('computer science and design', 'csd'):
                att_query = att_query.filter(func.lower(User.department).in_(['computer science and design', 'csd']))
            else:
                att_query = att_query.filter(func.lower(User.department) == dept_lower)
        avg_att = att_query.scalar()
        dept_avg_att = round(float(avg_att), 1) if avg_att is not None else 88.5

    return jsonify({
        'success': True,
        'stats': {
            'totalRequests': total,
            'totalODRequests': total,
            'totalODsSanctioned': approved,
            'pending': pending_hod,
            'pendingHOD': pending_hod,
            'pendingMentor': pending_mentor,
            'pendingClassIncharge': pending_ci,
            'approved': approved,
            'rejected': rejected,
            'approvalRate': round((approved / total * 100), 1) if total > 0 else 0,
            'departmentAvgAttendance': dept_avg_att,
            'eventDistribution': event_dist
        }
    }), 200

@hod_bp.route('/od-requests/<request_id>/letter', methods=['GET'])
@role_required('HOD')
def view_hod_od_letter(request_id):
    """
    Securely download/view the uploaded OD letter for an assigned request.
    Prevents unauthorized access and directory traversal.
    """
    hod_user = current_user
    req = ODRequestModel.get_by_id(request_id)

    if not req:
        return jsonify({'success': False, 'error': 'OD Request not found.'}), 404

    if not is_hod_authorized_for_request(hod_user, req):
        return jsonify({'success': False, 'error': 'Access denied.'}), 403

    doc_url = req.get('od_letter_url') or req.get('documentUrl')
    if not doc_url:
        return jsonify({'success': False, 'error': 'No OD letter attached to this request.'}), 404

    if doc_url.startswith('http://') or doc_url.startswith('https://'):
        return redirect(doc_url)

    filename = os.path.basename(doc_url)
    safe_filename = secure_filename(filename)

    file_path = os.path.join(Config.OD_LETTERS_FOLDER, safe_filename)
    if not os.path.exists(file_path):
        return jsonify({'success': False, 'error': 'Attached OD letter file not found on server.'}), 404

    return send_from_directory(Config.OD_LETTERS_FOLDER, safe_filename)


# ─── HOD DEPARTMENT FACULTY MANAGEMENT ────────────────────────────────────────

@hod_bp.route('/faculty', methods=['GET'])
@role_required('HOD')
def get_hod_department_faculty():
    """Retrieve faculty members within the HOD's department oversight."""
    hod_user = current_user
    hod_dept = (hod_user.get('department') or '').strip().lower()

    with session_scope() as session:
        query = session.query(User).filter(
            func.lower(User.role).in_(['mentor', 'class incharge', 'class_incharge', 'faculty', 'hod'])
        )
        if hod_dept and hod_dept != 'all':
            query = query.filter(func.lower(User.department) == hod_dept)

        faculty_records = query.order_by(User.name.asc()).all()
        
        # Count students in the department to report ward/assigned metrics
        student_count = session.query(User).filter(
            or_(func.lower(User.role) == 'student', func.lower(User.sub_role) == 'student'),
            func.lower(User.department) == hod_dept if (hod_dept and hod_dept != 'all') else True
        ).count()

        faculty_list = []
        for f in faculty_records:
            d = f.to_safe_dict()
            faculty_list.append({
                'id': d.get('id'),
                'faculty_id': d.get('id'),
                'employee_id': d.get('identifier') or d.get('id'),
                'name': d.get('name') or 'Faculty Member',
                'email': d.get('email') or '',
                'role': d.get('role') or 'Mentor',
                'sub_role': d.get('sub_role') or d.get('role'),
                'department': d.get('department'),
                'designation': d.get('designation') or 'Assistant Professor',
                'assigned_section': d.get('section') or 'All Sections',
                'phone': d.get('phone') or '',
                'assignedStudents': student_count if d.get('role') in ['Mentor', 'Class Incharge'] else 0,
                'status': d.get('account_status', 'ACTIVE'),
                'avatar': d.get('avatar'),
                'created_at': d.get('created_at')
            })

    return jsonify({
        'success': True,
        'faculty': faculty_list,
        'count': len(faculty_list)
    }), 200


# ─── HOD DEPARTMENT STUDENTS ──────────────────────────────────────────────────

@hod_bp.route('/students', methods=['GET'])
@role_required('HOD')
def get_hod_department_students():
    """Retrieve all students enrolled under HOD's department oversight."""
    hod_user = current_user
    hod_dept = (hod_user.get('department') or '').strip()

    students = UserModel.list_all(role='Student', department=hod_dept if (hod_dept and hod_dept.lower() != 'all') else None)

    acad_docs = {a['student_id']: a for a in AcademicRecordModel.list_all() if a and 'student_id' in a}

    roster = []
    for s in students:
        sid = s.get('id')
        reg_no = s.get('identifier') or sid
        acad = acad_docs.get(sid) or AcademicRecordModel.get_by_register_number(reg_no)

        elig = calculate_od_eligibility(sid, requested_od_days=0.0)

        att = acad.get('attendance_percentage') if acad else s.get('attendance_percentage', 85.0)
        cgpa = acad.get('cgpa') if acad and acad.get('cgpa') is not None else s.get('cgpa', 8.5)
        c1 = acad.get('cat1_marks') if acad else s.get('cat1_marks')
        c2 = acad.get('cat2_marks') if acad else s.get('cat2_marks')
        c3 = acad.get('cat3_marks') if acad else s.get('cat3_marks')

        roster.append({
            'id': sid,
            'studentId': sid,
            'registerNumber': reg_no,
            'rollNo': reg_no,
            'name': s.get('name') or 'Student',
            'studentName': s.get('name') or 'Student',
            'email': s.get('email') or '',
            'department': s.get('department'),
            'year': s.get('year'),
            'section': s.get('section', 'A'),
            'phone': s.get('phone', ''),
            'avatar': s.get('avatar'),
            'attendancePercent': float(att) if att is not None else 85.0,
            'cgpa': float(cgpa) if cgpa is not None else 8.5,
            'cat1Marks': c1,
            'cat2Marks': c2,
            'cat3Marks': c3,
            'odUsedDays': elig.get('od_used_days', 0.0),
            'eligibilityStatus': elig.get('eligibility_status_label', 'Eligible')
        })

    return jsonify({
        'success': True,
        'students': roster,
        'count': len(roster)
    }), 200


# ─── HOD CERTIFICATE SUBMISSIONS & AUDIT ───────────────────────────────────────

@hod_bp.route('/certificates', methods=['GET'])
@role_required('HOD')
def get_hod_certificates():
    """Retrieve certificates submitted by students within HOD's department."""
    hod_user = current_user
    hod_dept = (hod_user.get('department') or '').strip().lower()

    with session_scope() as session:
        cert_query = session.query(Certificate).join(ODRequest, Certificate.request_id == ODRequest.id)
        if hod_dept and hod_dept != 'all':
            cert_query = cert_query.filter(func.lower(ODRequest.department) == hod_dept)
        certs = cert_query.order_by(Certificate.created_at.desc()).all()
        cert_list = [c.to_dict() for c in certs]

        # Also include OD requests that have certificate_url but no entry in certificates table yet
        od_reqs = session.query(ODRequest).filter(
            ODRequest.certificate_url.isnot(None),
            ODRequest.certificate_url != ''
        )
        if hod_dept and hod_dept != 'all':
            od_reqs = od_reqs.filter(func.lower(ODRequest.department) == hod_dept)

        existing_req_ids = {c.get('request_id') for c in cert_list}
        for req in od_reqs.all():
            if req.id not in existing_req_ids:
                fname = os.path.basename(req.certificate_url or 'certificate.pdf')
                cert_list.append({
                    'id': f"CERT_{req.id}",
                    'requestId': req.id,
                    'request_id': req.id,
                    'studentId': req.student_id,
                    'student_id': req.student_id,
                    'studentRegNo': req.student_reg_no,
                    'student_reg_no': req.student_reg_no,
                    'studentName': req.student_name,
                    'student_name': req.student_name,
                    'eventName': req.event_name,
                    'event_name': req.event_name,
                    'certificateName': fname,
                    'certificate_file_url': req.certificate_url,
                    'certificateUrl': req.certificate_url,
                    'status': req.certificate_status or 'Pending Verification',
                    'uploadDate': to_iso(req.certificate_submitted_at or req.updated_at),
                    'upload_date': to_iso(req.certificate_submitted_at or req.updated_at),
                })

    return jsonify({'success': True, 'certificates': cert_list, 'count': len(cert_list)}), 200


@hod_bp.route('/od-requests/<request_id>/verify-certificate', methods=['POST'])
@role_required('HOD')
def verify_student_certificate_hod(request_id):
    """
    Verify student participation certificate as Head of Department (HOD).
    Status can be 'Verified' or 'Rejected'.
    """
    hod_user = current_user
    data = request.get_json(silent=True) or {}
    status = (data.get('status') or 'Verified').strip()
    remarks = (data.get('remarks') or '').strip()

    updated_req, err = ODRequestModel.verify_certificate(request_id, hod_user, status, remarks)
    if err:
        return jsonify({'success': False, 'error': err}), 400

    return jsonify({
        'success': True,
        'message': f"Certificate successfully {status.lower()} by HOD {hod_user.get('name')}.",
        'request': updated_req
    }), 200


