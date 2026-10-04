import os
import sys
import io
import openpyxl
import mongomock
from datetime import datetime
from werkzeug.security import generate_password_hash

# Ensure root is in sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from backend.database.mongodb import MongoDB
from backend.models.od_request import ODRequestModel
from backend.app import create_app

def setup_test_data():
    mock_client = mongomock.MongoClient()
    MongoDB._client = mock_client
    MongoDB._db = mock_client['od_tracking']
    MongoDB._is_connected = True
    MongoDB._connection_status = "connected"
    MongoDB._connection_error = None

    db = MongoDB._db
    default_pwd_hash = generate_password_hash('password123')

    # Seed users
    db['users'].insert_many([
        {
            'id': 'STUD001',
            'identifier': '23CSD001',
            'name': 'Naveen',
            'email': 'naveen.23csd@rajalakshmi.edu.in',
            'role': 'Student',
            'sub_role': 'Student',
            'department': 'Computer Science and Design',
            'year': 'III Year',
            'section': 'A',
            'designation': 'Student',
            'password_hash': default_pwd_hash,
            'created_at': '2026-01-01 09:00:00',
            'updated_at': '2026-01-01 09:00:00'
        },
        {
            'id': 'STUD002',
            'identifier': '23CSD002',
            'name': 'Priya S',
            'email': 'priya.23csd@rajalakshmi.edu.in',
            'role': 'Student',
            'sub_role': 'Student',
            'department': 'Computer Science and Design',
            'year': 'III Year',
            'section': 'A',
            'designation': 'Student',
            'password_hash': default_pwd_hash,
            'created_at': '2026-01-01 09:00:00',
            'updated_at': '2026-01-01 09:00:00'
        },
        {
            'id': 'FAC004',
            'identifier': 'EMP-CSD-104',
            'name': 'Dr. V. Karpagam',
            'email': 'v.karpagam@rajalakshmi.edu.in',
            'role': 'HOD',
            'sub_role': 'HOD',
            'department': 'Computer Science and Design',
            'year': 'All Years',
            'section': 'All Sections (CSD)',
            'designation': 'Professor & HOD',
            'password_hash': default_pwd_hash,
            'created_at': '2026-01-01 09:00:00',
            'updated_at': '2026-01-01 09:00:00'
        }
    ])

    # Seed certificates
    db['certificates'].insert_many([
        {
            'id': 'CERT_20260901_001',
            'request_id': 'REQ_20260901_001',
            'student_id': 'STUD001',
            'student_reg_no': '23CSD001',
            'event_name': 'Smart India Hackathon 2026',
            'certificate_file_url': '/api/od-requests/uploads/certificates/sih2026.pdf',
            'upload_date': '2026-09-13 10:00:00',
            'status': 'Verified',
            'verified_by_id': 'FAC001',
            'verified_by_name': 'Dr. A. Rajesh',
            'verified_at': '2026-09-14 11:00:00',
            'remarks': '1st Prize Certificate verified with national portal.',
            'created_at': '2026-09-13 10:00:00'
        }
    ])

    # Seed diverse OD Requests including minimal student_id reference and deleted student
    db['od_requests'].insert_many([
        {
            'id': 'REQ_20260901_001',
            'student_id': 'STUD001',
            'student_reg_no': '23CSD001',
            'department': 'Computer Science and Design',
            'event_name': 'Smart India Hackathon 2026',
            'event_type': 'Hackathon',
            'event_organizer': 'AICTE & MoE',
            'venue': 'IIT Madras Research Park',
            'from_date': '2026-09-10',
            'to_date': '2026-09-12',
            'from_time': '09:00',
            'to_time': '18:00',
            'number_of_days': 3.0,
            'reason': 'Grand Finale Hackathon',
            'description': 'Finalist team presentation',
            'od_letter_url': '/api/od-requests/uploads/od_letters/sample.pdf',
            'status': 'Approved',
            'current_stage': 'Approved',
            'certificate_status': 'Verified',
            'remarks': 'Approved by HOD with full attendance credit',
            'created_at': '2026-09-01 10:00:00',
            'updated_at': '2026-09-05 14:00:00'
        },
        {
            'id': 'REQ_20260902_002',
            'student_id': 'STUD002',
            'student_reg_no': '23CSD002',
            'department': 'Computer Science and Design',
            'event_name': 'National Design Symposium',
            'event_type': 'Symposium',
            'event_organizer': 'NIT Trichy',
            'venue': 'NIT Trichy Campus',
            'from_date': '2026-09-18',
            'to_date': '2026-09-19',
            'from_time': '09:30',
            'to_time': '17:00',
            'number_of_days': 2.0,
            'reason': 'Paper Presentation on UI/UX in AI',
            'description': 'Paper accepted for oral presentation',
            'od_letter_url': '/api/od-requests/uploads/od_letters/symposium.pdf',
            'status': 'Class Incharge Approved',
            'current_stage': 'HOD',
            'certificate_status': 'Not Uploaded',
            'remarks': 'Endorsed by Class Incharge',
            'created_at': '2026-09-02 11:30:00',
            'updated_at': '2026-09-04 16:00:00'
        },
        {
            'id': 'REQ_20260903_003',
            'student_id': 'STUD001',
            'student_reg_no': '23CSD001',
            'department': 'Computer Science and Design',
            'event_name': 'State Level Badminton Championship',
            'event_type': 'Sports',
            'event_organizer': 'Anna University Sports Board',
            'venue': 'Anna University Stadium, Chennai',
            'from_date': '2026-09-25',
            'to_date': '2026-09-26',
            'from_time': '08:00',
            'to_time': '17:00',
            'number_of_days': 2.0,
            'reason': 'Representing College Badminton Team',
            'description': 'Selected for Inter-Zonal round',
            'od_letter_url': '/api/od-requests/uploads/od_letters/sports.pdf',
            'status': 'Approved',
            'current_stage': 'Approved',
            'certificate_status': 'Pending Verification',
            'remarks': 'Approved by HOD',
            'created_at': '2026-09-03 09:15:00',
            'updated_at': '2026-09-06 11:00:00'
        },
        {
            'id': 'REQ_20260904_004',
            'student_id': 'STUD002',
            'student_reg_no': '23CSD002',
            'department': 'Computer Science and Design',
            'event_name': 'Web3 Developer Bootcamp',
            'event_type': 'Workshop',
            'event_organizer': 'Ethereum India',
            'venue': 'Online / Virtual',
            'from_date': '2026-10-05',
            'to_date': '2026-10-06',
            'from_time': '10:00',
            'to_time': '16:00',
            'number_of_days': 2.0,
            'reason': 'Online Web3 Workshop',
            'description': 'Virtual hands-on lab',
            'od_letter_url': None,
            'status': 'HOD Rejected',
            'current_stage': 'HOD',
            'certificate_status': 'Not Uploaded',
            'remarks': 'Declined: Clashes with internal assessment tests.',
            'created_at': '2026-09-04 14:00:00',
            'updated_at': '2026-09-07 10:30:00'
        },
        # Record with ONLY student_id (no embedded student_name, student_reg_no, etc.)
        {
            'id': 'REQ_20260905_005',
            'student_id': 'STUD001',
            'department': 'Computer Science and Design',
            'event_name': 'Google Cloud Community Day',
            'event_type': 'Conference',
            'event_organizer': 'GDG Chennai',
            'venue': 'Chennai Trade Centre',
            'from_date': '2026-10-12',
            'to_date': '2026-10-12',
            'from_time': '09:00',
            'to_time': '17:00',
            'number_of_days': 1.0,
            'reason': 'Tech Conference',
            'description': 'Cloud architecture tracks',
            'od_letter_url': None,
            'status': 'Approved',
            'current_stage': 'Approved',
            'certificate_status': 'Not Uploaded',
            'remarks': 'Approved by HOD',
            'created_at': '2026-09-05 08:30:00',
            'updated_at': '2026-09-08 12:00:00'
        },
        # Edge case record: Non-existent / deleted student
        {
            'id': 'REQ_20260906_006',
            'student_id': 'STUD_DELETED_999',
            'department': 'Computer Science and Design',
            'event_name': 'Inter-College AI Hackfest',
            'event_type': 'Hackathon',
            'event_organizer': 'Tech Club',
            'venue': 'Auditorium',
            'from_date': '2026-10-20',
            'to_date': '2026-10-20',
            'from_time': '09:00',
            'to_time': '17:00',
            'number_of_days': 1.0,
            'reason': 'AI Coding',
            'description': 'Competitive programming',
            'od_letter_url': None,
            'status': 'Approved',
            'current_stage': 'Approved',
            'certificate_status': 'Not Uploaded',
            'remarks': 'Approved by HOD',
            'created_at': '2026-09-06 09:00:00',
            'updated_at': '2026-09-08 13:00:00'
        }
    ])

def test_hod_excel_export():
    print("=== TESTING HOD EXCEL EXPORT INTEGRATION ===")
    setup_test_data()
    app = create_app()

    with app.test_client() as client:
        # 1. Login as HOD
        res = client.post('/api/auth/login', json={'identifier': 'EMP-CSD-104', 'password': 'password123'})
        assert res.status_code == 200
        hod_token = res.get_json()['token']
        print("[1] HOD Login: Success (Token acquired)")

        # 2. Login as Student
        res = client.post('/api/auth/login', json={'identifier': '23CSD001', 'password': 'password123'})
        assert res.status_code == 200
        student_token = res.get_json()['token']

        # 3. Test Security: Student cannot access HOD Export endpoint
        res = client.get('/api/hod/od/export', headers={'Authorization': f'Bearer {student_token}'})
        print(f"[2] Security Check: Student export attempt blocked -> HTTP {res.status_code}")
        assert res.status_code == 403

        # 4. Test Export: All Records
        res = client.get('/api/hod/od/export', headers={'Authorization': f'Bearer {hod_token}'})
        print(f"[3] General Export All Records -> HTTP {res.status_code}, Content-Type: {res.content_type}")
        assert res.status_code == 200
        assert "spreadsheetml" in res.content_type

        # Verify Excel binary content
        excel_bytes = res.data
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active
        assert ws.title == "OD Requests Report"
        print(f"  - Workbook loaded successfully: {ws.max_row} rows x {ws.max_column} columns")
        
        # Verify header title
        assert "RAJALAKSHMI ENGINEERING COLLEGE" in str(ws.cell(row=1, column=1).value)
        # Verify table headers in row 4
        assert ws.cell(row=4, column=1).value == "S.No"
        assert ws.cell(row=4, column=2).value == "Request ID"
        assert ws.cell(row=4, column=3).value == "Student Name"
        assert ws.cell(row=4, column=4).value == "Register Number"
        assert ws.cell(row=4, column=5).value == "Department"
        assert ws.cell(row=4, column=6).value == "Year"
        assert ws.cell(row=4, column=7).value == "Section"
        assert ws.cell(row=4, column=8).value == "Event / Program Name"
        assert ws.cell(row=4, column=24).value == "Certificate Remarks"
        assert ws.cell(row=4, column=25).value == "Final Completion Status"

        # Verify data rows
        rows_data = []
        for r in range(5, ws.max_row):
            req_id = ws.cell(row=r, column=2).value
            if not req_id or not str(req_id).startswith("REQ_"):
                continue
            row_dict = {
                "s_no": ws.cell(row=r, column=1).value,
                "req_id": ws.cell(row=r, column=2).value,
                "student_name": ws.cell(row=r, column=3).value,
                "reg_no": ws.cell(row=r, column=4).value,
                "department": ws.cell(row=r, column=5).value,
                "year": ws.cell(row=r, column=6).value,
                "section": ws.cell(row=r, column=7).value,
                "event_name": ws.cell(row=r, column=8).value,
                "event_type": ws.cell(row=r, column=9).value,
                "event_date": ws.cell(row=r, column=10).value,
                "venue": ws.cell(row=r, column=11).value,
                "reason": ws.cell(row=r, column=12).value,
                "request_date": ws.cell(row=r, column=13).value,
                "document": ws.cell(row=r, column=14).value,
                "mentor_status": ws.cell(row=r, column=15).value,
                "mentor_remarks": ws.cell(row=r, column=16).value,
                "ci_status": ws.cell(row=r, column=17).value,
                "ci_remarks": ws.cell(row=r, column=18).value,
                "hod_status": ws.cell(row=r, column=19).value,
                "hod_remarks": ws.cell(row=r, column=20).value,
                "overall_status": ws.cell(row=r, column=21).value,
                "cert_uploaded": ws.cell(row=r, column=22).value,
                "cert_status": ws.cell(row=r, column=23).value,
                "cert_remarks": ws.cell(row=r, column=24).value,
                "final_status": ws.cell(row=r, column=25).value
            }
            rows_data.append(row_dict)

        print(f"  - Extracted {len(rows_data)} data rows from spreadsheet.")
        
        # Verify student actual data is populated for Naveen (REQ_20260901_001)
        sih_row = next(row for row in rows_data if row['req_id'] == 'REQ_20260901_001')
        print(f"  - Row REQ_20260901_001: {sih_row}")
        assert sih_row['student_name'] == 'Naveen'
        assert sih_row['reg_no'] == '23CSD001'
        assert sih_row['department'] == 'Computer Science and Design'
        assert sih_row['year'] == 'III Year'
        assert sih_row['section'] == 'A'
        assert sih_row['event_name'] == 'Smart India Hackathon 2026'
        assert sih_row['overall_status'] == 'Approved'
        assert sih_row['cert_uploaded'] == 'Yes'
        assert sih_row['cert_status'] == 'Verified'
        assert '1st Prize Certificate' in sih_row['cert_remarks']

        # Verify dynamic join for minimal record (REQ_20260905_005: only had student_id)
        minimal_row = next(row for row in rows_data if row['req_id'] == 'REQ_20260905_005')
        print(f"  - Row REQ_20260905_005 (minimal OD request joined with users collection): {minimal_row}")
        assert minimal_row['student_name'] == 'Naveen'
        assert minimal_row['reg_no'] == '23CSD001'
        assert minimal_row['department'] == 'Computer Science and Design'
        assert minimal_row['year'] == 'III Year'
        assert minimal_row['section'] == 'A'

        # Verify edge case for deleted/non-existent student (REQ_20260906_006)
        deleted_row = next(row for row in rows_data if row['req_id'] == 'REQ_20260906_006')
        print(f"  - Row REQ_20260906_006 (deleted student edge case): {deleted_row}")
        assert deleted_row['student_name'] == 'Student Not Found'
        assert deleted_row['reg_no'] == 'N/A'

        # 5. Test Filter: Event Type = Hackathon
        res = client.get('/api/hod/od/export?event_type=Hackathon', headers={'Authorization': f'Bearer {hod_token}'})
        assert res.status_code == 200
        wb_hack = openpyxl.load_workbook(io.BytesIO(res.data))
        ws_hack = wb_hack.active
        hack_events = [ws_hack.cell(row=r, column=8).value for r in range(5, ws_hack.max_row) if ws_hack.cell(row=r, column=8).value and ws_hack.cell(row=r, column=2).value]
        print(f"[4] Filter Event Type=Hackathon -> Exported Events: {hack_events}")
        assert "Smart India Hackathon 2026" in hack_events
        assert "State Level Badminton Championship" not in hack_events

        # 6. Test Filter: Status = Approved
        res = client.get('/api/hod/od/export?status=Approved', headers={'Authorization': f'Bearer {hod_token}'})
        assert res.status_code == 200
        wb_app = openpyxl.load_workbook(io.BytesIO(res.data))
        ws_app = wb_app.active
        app_statuses = [ws_app.cell(row=r, column=21).value for r in range(5, ws_app.max_row) if ws_app.cell(row=r, column=21).value and ws_app.cell(row=r, column=2).value]
        print(f"[5] Filter Status=Approved -> Found statuses: {app_statuses}")
        assert all(s == "Approved" for s in app_statuses)

        # 7. Test Filter: Date Range (2026-09-01 to 2026-09-15)
        res = client.get('/api/hod/od/export?from_date=2026-09-01&to_date=2026-09-15', headers={'Authorization': f'Bearer {hod_token}'})
        assert res.status_code == 200
        wb_date = openpyxl.load_workbook(io.BytesIO(res.data))
        ws_date = wb_date.active
        date_events = [ws_date.cell(row=r, column=8).value for r in range(5, ws_date.max_row) if ws_date.cell(row=r, column=8).value and ws_date.cell(row=r, column=2).value]
        print(f"[6] Filter Date Range (Sep 1 - Sep 15) -> Events: {date_events}")
        assert "Smart India Hackathon 2026" in date_events
        assert "Web3 Developer Bootcamp" not in date_events

        # 8. Test Filename in Content-Disposition Header
        disposition = res.headers.get('Content-Disposition')
        print(f"[7] Content-Disposition header: {disposition}")
        assert "attachment" in disposition and "filename=" in disposition

    print("\n========================================================")
    print("ALL HOD EXCEL EXPORT INTEGRATION TESTS PASSED 100%!")
    print("========================================================")

if __name__ == '__main__':
    test_hod_excel_export()
