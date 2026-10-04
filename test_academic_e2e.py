"""
End-to-End Academic Data Upload & Verification Test Script using standard urllib
"""
import io
import json
import urllib.request
import urllib.error
import openpyxl

BASE_URL = "http://127.0.0.1:5000"

def http_req(url, method="GET", data=None, headers=None):
    if headers is None:
        headers = {}
    body = None
    if data is not None:
        if isinstance(data, dict):
            body = json.dumps(data).encode("utf-8")
            headers["Content-Type"] = "application/json"
        elif isinstance(data, (bytes, bytearray)):
            body = data
    req = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req) as resp:
            content = resp.read()
            return resp.status, content, resp.headers
    except urllib.error.HTTPError as e:
        content = e.read()
        return e.code, content, e.headers

def multipart_upload(url, filename, file_bytes, headers=None):
    boundary = "----WebKitFormBoundary7MA4YWxkTrZu0gW"
    if headers is None:
        headers = {}
    headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"
    
    parts = []
    parts.append(f"--{boundary}\r\n".encode())
    parts.append(f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'.encode())
    parts.append(b"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n")
    parts.append(file_bytes)
    parts.append(f"\r\n--{boundary}--\r\n".encode())
    
    payload = b"".join(parts)
    return http_req(url, method="POST", data=payload, headers=headers)

def run_tests():
    print("=== STARTING ACADEMIC MODULE E2E TESTS ===")

    # 1. Login as Class Incharge
    status, body, _ = http_req(
        f"{BASE_URL}/api/auth/login",
        method="POST",
        data={"identifier": "EMP-CSD-102", "password": "password123", "role": "Class Incharge"}
    )
    assert status == 200, f"Login failed: {body.decode()}"
    ci_token = json.loads(body.decode()).get("token")
    headers_ci = {"Authorization": f"Bearer {ci_token}"}
    print("[PASS] 1. Class Incharge logged in successfully.")

    # 2. Download sample Excel template
    status, body, resp_hdrs = http_req(
        f"{BASE_URL}/api/class-incharge/academic/template",
        headers=headers_ci
    )
    assert status == 200, f"Template download failed: {status}"
    
    wb = openpyxl.load_workbook(io.BytesIO(body))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    expected_headers = ['Register Number', 'Student Name', 'CAT 1', 'CAT 2', 'CAT 3', 'Attendance Percentage']
    assert headers == expected_headers, f"Template mismatch: {headers}"
    print("[PASS] 2. Downloaded Excel template verified with exact columns.")

    # 3. Test Invalid Excel Validation
    invalid_wb = openpyxl.Workbook()
    invalid_ws = invalid_wb.active
    invalid_ws.append(expected_headers)
    invalid_ws.append(['23CSD001', 'Naveen S', -5, 82, 75, 91]) # invalid CAT 1 < 0
    invalid_ws.append(['23CSD002', 'Priya S', 65, 70, 68, 150]) # invalid Attendance > 100
    invalid_buf = io.BytesIO()
    invalid_wb.save(invalid_buf)

    status, body, _ = multipart_upload(
        f"{BASE_URL}/api/class-incharge/academic/upload",
        "invalid_test.xlsx",
        invalid_buf.getvalue(),
        headers=headers_ci
    )
    assert status == 200
    invalid_json = json.loads(body.decode())
    assert invalid_json.get("summary", {}).get("can_confirm") is False
    assert len(invalid_json.get("errors", [])) == 2
    print(f"[PASS] 3. Data validation passed: {len(invalid_json.get('errors'))} invalid rows flagged, update blocked.")

    # 4. Test Valid Excel Upload
    valid_wb = openpyxl.Workbook()
    valid_ws = valid_wb.active
    valid_ws.append(expected_headers)
    valid_ws.append(['23CSD001', 'Naveen S', 88, 92, 85, 94]) # Existing Student 1
    valid_ws.append(['23CSD002', 'Priya S', 72, 78, 80, 89])  # Existing Student 2
    valid_ws.append(['23CSD003', 'Karthik R', 68, 74, 70, 82]) # Existing Student 3
    valid_ws.append(['241701999', 'Ghost Student', 80, 85, 75, 90]) # Unmatched Student
    valid_buf = io.BytesIO()
    valid_wb.save(valid_buf)

    status, body, _ = multipart_upload(
        f"{BASE_URL}/api/class-incharge/academic/upload",
        "CSD_CAT_Attendance_Oct.xlsx",
        valid_buf.getvalue(),
        headers=headers_ci
    )
    assert status == 200
    valid_json = json.loads(body.decode())
    assert valid_json.get("summary", {}).get("can_confirm") is True
    summary = valid_json.get("summary", {})
    assert summary.get("total_rows") == 4
    assert summary.get("matched_count") == 3
    assert summary.get("unmatched_count") == 1
    assert summary.get("invalid_count") == 0
    confirm_token = valid_json.get("confirm_token")
    assert confirm_token
    print("[PASS] 4. Valid Excel preview generated: 3 matched, 1 unmatched, 0 invalid rows.")

    # 5. Confirm and update database
    status, body, _ = http_req(
        f"{BASE_URL}/api/class-incharge/academic/confirm",
        method="POST",
        data={"confirm_token": confirm_token},
        headers=headers_ci
    )
    assert status == 200
    confirm_json = json.loads(body.decode())
    assert confirm_json.get("success") is True
    assert confirm_json.get("updated_count") == 3
    print(f"[PASS] 5. Confirmed database update: {confirm_json.get('updated_count')} records updated.")

    # 6. Verify Upload History Log
    status, body, _ = http_req(
        f"{BASE_URL}/api/class-incharge/academic/upload-history",
        headers=headers_ci
    )
    assert status == 200
    history_list = json.loads(body.decode()).get("history", [])
    assert len(history_list) >= 1
    latest_hist = history_list[0]
    assert latest_hist.get("file_name") == "CSD_CAT_Attendance_Oct.xlsx"
    assert latest_hist.get("successful_updates") == 3
    assert latest_hist.get("unmatched_students") == 1
    print("[PASS] 6. Upload audit history verified.")

    # 7. Login as Student (23CSD001) and verify live attendance & CAT marks
    status, body, _ = http_req(
        f"{BASE_URL}/api/auth/login",
        method="POST",
        data={"identifier": "23CSD001", "password": "password123", "role": "Student"}
    )
    assert status == 200
    student_token = json.loads(body.decode()).get("token")
    headers_student = {"Authorization": f"Bearer {student_token}"}

    status, body, _ = http_req(
        f"{BASE_URL}/api/academic/me",
        headers=headers_student
    )
    assert status == 200
    acad_data = json.loads(body.decode()).get("academic", {})
    assert acad_data.get("attendance_percentage") == 94, f"Attendance mismatch: {acad_data}"
    assert acad_data.get("cat1_marks") == 88
    assert acad_data.get("cat2_marks") == 92
    assert acad_data.get("cat3_marks") == 85
    
    eligibility = acad_data.get("eligibility", {})
    assert round(eligibility.get("max_od_allowed_days")) == 11, f"Eligibility: {eligibility}"
    print(f"[PASS] 7. Student view verified: Attendance = 94%, CAT1=88, CAT2=92, CAT3=85, 10% OD Cap = {eligibility.get('max_od_allowed_days')} days.")

    # 8. Check Student cannot upload (Role Permissions)
    status, body, _ = multipart_upload(
        f"{BASE_URL}/api/class-incharge/academic/upload",
        "hack.xlsx",
        valid_buf.getvalue(),
        headers=headers_student
    )
    assert status == 403, f"Expected 403 Forbidden, got {status}"
    print("[PASS] 8. Role-based security enforced: Student upload attempt blocked (403 Forbidden).")

    # 9. Mentor/Faculty single student inspection
    status, body, _ = http_req(
        f"{BASE_URL}/api/auth/login",
        method="POST",
        data={"identifier": "EMP-CSD-101", "password": "password123", "role": "Mentor"}
    )
    assert status == 200, f"Mentor login failed: {body.decode()}"
    mentor_token = json.loads(body.decode()).get("token")
    status, body, _ = http_req(
        f"{BASE_URL}/api/students/23CSD001/academic",
        headers={"Authorization": f"Bearer {mentor_token}"}
    )
    assert status == 200
    inspect_data = json.loads(body.decode()).get("academic", {})
    assert inspect_data.get("attendance_percentage") == 94
    assert inspect_data.get("cat1_marks") == 88
    print("[PASS] 9. Mentor/Faculty inspection verified: Attendance = 94%.")

    print("\n===========================================")
    print("ALL 9 E2E ACADEMIC TESTS PASSED PERFECTLY!")
    print("===========================================")

if __name__ == "__main__":
    run_tests()
